"""Erzeugt die Excel-Vorlage "Lademeter berechnen" (ohne Makros).

Aufruf:  python seo/tools/build_lademeter_excel.py [Ausgabedatei]

Die Formeln bilden exakt die Rechenkerne von lademeter-rechner.de nach (lib/calculators/lademeter.ts):
  Bodenstellplaetze = ROUNDUP(Anzahl / Stapellagen)
  Lademeter         = Laenge[m] * Breite[m] / Referenzbreite * Bodenstellplaetze
  frachtpflichtiges Gewicht = MAX(tats. Gewicht, Lademeter * Faktor)
Verifiziert durch seo/tools/verify_lademeter_excel.py.
"""
import datetime
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

OUT = sys.argv[1] if len(sys.argv) > 1 else "seo/assets/Lademeter-Berechnung-Excel-Vorlage.xlsx"
VERSION = "1.0"
STAND = "20.10.2026"
SITE = "https://lademeter-rechner.de"

NAVY = "0B1324"
ORANGE = "F97316"
INPUT_FILL = PatternFill("solid", fgColor="EAF2FF")
CALC_FILL = PatternFill("solid", fgColor="F1F5F9")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
KEY_FILL = PatternFill("solid", fgColor="FFF1E6")
WARN_FILL = PatternFill("solid", fgColor="FEE2E2")
thin = Side(style="thin", color="CBD5E1")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
F_BASE = Font(name="Calibri", size=11, color="0B1324")
F_BOLD = Font(name="Calibri", size=11, bold=True, color="0B1324")
F_INPUT = Font(name="Calibri", size=11, color="1D4ED8")
F_HEAD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
F_TITLE = Font(name="Calibri", size=18, bold=True, color="0B1324")
F_SMALL = Font(name="Calibri", size=9, color="64748B")
F_LINK = Font(name="Calibri", size=11, underline="single", color="C2410C")

FIRST, LAST = 14, 38  # Datenzeilen (25 Positionen)
TOTAL = LAST + 1

EXAMPLE_ROWS = [
    ("Europaletten", 120, 80, 150, 10, 1, 420),
    ("Industriepaletten (2 Lagen)", 120, 100, 140, 6, 2, 600),
    ("Kisten", 85, 65, 90, 20, 1, 85),
]


def style_range(ws, ref, font=None, fill=None, border=None, align=None, number_format=None, unlocked=False):
    for row in ws[ref]:
        for cell in row:
            if font: cell.font = font
            if fill: cell.fill = fill
            if border: cell.border = border
            if align: cell.alignment = align
            if number_format: cell.number_format = number_format
            if unlocked: cell.protection = Protection(locked=False)


def build_calc_sheet(ws, example: bool):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ORANGE if not example else "64748B"
    widths = {"A": 5, "B": 40, "C": 12, "D": 12, "E": 12, "F": 10, "G": 12, "H": 15, "I": 16, "J": 14, "K": 14, "L": 14, "M": 15, "N": 44}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws["A1"] = "Lademeter berechnen – Excel-Vorlage" + (" (Beispiel)" if example else "")
    ws["A1"].font = F_TITLE
    ws["A2"] = (
        "Blau hinterlegte Zellen sind Eingaben, graue Zellen werden berechnet. Rechenweg und Hinweise: Blatt „Anleitung“. "
        f"Online-Rechner und Erklärungen: {SITE}"
    )
    ws["A2"].font = F_SMALL

    # --- Einstellungen
    ws["A3"] = "Einstellungen"
    ws["A3"].font = F_BOLD
    labels = [
        (4, "Maßeinheit der Eingaben (mm, cm, m)", "cm"),
        (5, "Referenzbreite der Ladefläche (m)", 2.4),
        (6, "Umrechnungsfaktor kg/ldm (optional)", 1850 if example else None),
        (7, "Ladelänge Vergleichsfahrzeug (m)", 13.6),
    ]
    for r, text, value in labels:
        ws.cell(r, 2, text).font = F_BASE
        c = ws.cell(r, 3, value)
        c.fill, c.font, c.border = INPUT_FILL, F_INPUT, BORDER
        c.protection = Protection(locked=False)
        c.alignment = Alignment(horizontal="right")
    ws["C5"].number_format = "0.00"
    ws["C6"].number_format = "#,##0"
    ws["C7"].number_format = "0.0"
    ws["B8"] = "Hilfswert: Divisor zur Umrechnung in m"
    ws["B8"].font = F_SMALL
    ws["C8"] = '=IF(C4="mm",1000,IF(C4="cm",100,1))'
    ws["C8"].fill, ws["C8"].font, ws["C8"].border = CALC_FILL, F_SMALL, BORDER

    dv_unit = DataValidation(type="list", formula1='"mm,cm,m"', allow_blank=False, showErrorMessage=True,
                             errorTitle="Maßeinheit", error="Bitte mm, cm oder m wählen.")
    dv_pos = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True,
                            errorTitle="Ungültiger Wert", error="Bitte eine Zahl größer 0 eingeben.")
    dv_qty = DataValidation(type="whole", operator="between", formula1="1", formula2="99999", allow_blank=True,
                            showErrorMessage=True, errorTitle="Anzahl", error="Bitte eine ganze Zahl ab 1 eingeben.")
    dv_lay = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True,
                            showErrorMessage=True, errorTitle="Stapellagen",
                            error="Bitte 1 bis 10 eingeben (1 = nicht stapelbar).")
    for dv in (dv_unit, dv_pos, dv_qty, dv_lay):
        ws.add_data_validation(dv)
    dv_unit.add("C4")
    dv_pos.add("C5"); dv_pos.add("C6"); dv_pos.add("C7")
    dv_pos.add(f"C{FIRST}:E{LAST}"); dv_pos.add(f"H{FIRST}:H{LAST}")
    dv_qty.add(f"F{FIRST}:F{LAST}")
    dv_lay.add(f"G{FIRST}:G{LAST}")
    ws["C4"].value = "cm"

    # --- Ergebnis
    ws["E3"] = "Ergebnis"
    ws["E3"].font = F_BOLD
    rng = lambda col: f"{col}{FIRST}:{col}{LAST}"
    results = [
        (4, "E", "H", "Lademeter gesamt (ldm)", f"=ROUND(SUM({rng('J')}),2)", "0.00", True),
        (5, "E", "H", "Bodenstellplätze gesamt", f"=SUM({rng('I')})", "0", False),
        (6, "E", "H", "Grundfläche gesamt (m²)", f"=ROUND(SUM({rng('K')}),2)", "0.00", False),
        (7, "E", "H", "Volumen gesamt (m³) *", f"=ROUND(SUM({rng('L')}),2)", "0.00", False),
        (8, "E", "H", "Tatsächliches Gewicht (kg) *", f"=ROUND(SUM({rng('M')}),2)", "#,##0.00", False),
        (4, "J", "M", "Lademetergewicht (kg)",
         f'=IF(AND(ISNUMBER($C$6),$C$6>0,SUM({rng("J")})>0),ROUND(SUM({rng("J")})*$C$6,2),"–")', "#,##0.00", False),
        (5, "J", "M", "Frachtpflichtiges Gewicht (kg)",
         f'=IF(AND(ISNUMBER($C$6),$C$6>0,SUM({rng("J")})>0),ROUND(MAX(SUM({rng("M")}),SUM({rng("J")})*$C$6),2),"–")', "#,##0.00", True),
        (6, "J", "M", "Anteil Vergleichsfahrzeug (%)",
         f'=IF(SUM({rng("J")})>0,ROUND(SUM({rng("J")})/$C$7*100,2),"–")', "0.00", False),
        (7, "J", "M", "Rest-Lademeter (Vergleichsfzg.)",
         f'=IF(SUM({rng("J")})>0,ROUND($C$7-SUM({rng("J")}),2),"–")', "0.00", False),
    ]
    for r, c1, vcol, label, formula, fmt, key in results:
        ws[f"{c1}{r}"] = label
        ws[f"{c1}{r}"].font = F_BOLD if key else F_BASE
        cell = ws[f"{vcol}{r}"]
        cell.value = formula
        cell.number_format = fmt
        cell.font = Font(name="Calibri", size=12 if key else 11, bold=True, color="0B1324")
        cell.fill = KEY_FILL if key else CALC_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="right")
    ws["J8"] = f'=IF(SUM({rng("J")})=0,"Bitte Positionen eintragen.",IF(SUM({rng("J")})>$C$7,"Achtung: Die Sendung überschreitet die Ladelänge des Vergleichsfahrzeugs.","Die Sendung passt rechnerisch auf die Ladelänge des Vergleichsfahrzeugs (nur Fläche, keine Gewichts- oder Sicherungsprüfung)."))'
    ws["J8"].font = F_SMALL
    ws["E9"] = "* nur Positionen mit Höhen- bzw. Gewichtsangabe. Frachtpflichtiges Gewicht = höherer Wert aus tatsächlichem Gewicht und Lademeter × Faktor."
    ws["E9"].font = F_SMALL
    ws["J8"].alignment = Alignment(wrap_text=False)

    # --- Tabelle
    heads = ["Nr.", "Bezeichnung", "Länge", "Breite", "Höhe (optional)", "Anzahl", "Stapellagen", "Gewicht je Stück (kg, optional)",
             "Bodenstellplätze", "Lademeter (ldm)", "Grundfläche (m²)", "Volumen (m³)", "Gewicht gesamt (kg)", "Prüfhinweis"]
    ws["A12"] = "Ladungspositionen (bis zu 25)"
    ws["A12"].font = F_BOLD
    for i, h in enumerate(heads, start=1):
        c = ws.cell(13, i, h)
        c.fill, c.font, c.border = HEAD_FILL, F_HEAD, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[13].height = 44

    for k, r in enumerate(range(FIRST, LAST + 1), start=1):
        ws.cell(r, 1, k).font = F_SMALL
        ws.cell(r, 1).alignment = Alignment(horizontal="center")
        ws.cell(r, 1).border = BORDER
        data = EXAMPLE_ROWS[k - 1] if example and k <= len(EXAMPLE_ROWS) else None
        for col in range(2, 9):
            cell = ws.cell(r, col)
            if data is not None:
                cell.value = data[col - 2]
            cell.fill, cell.font, cell.border = INPUT_FILL, F_INPUT, BORDER
            cell.protection = Protection(locked=False)
            if col >= 3:
                cell.alignment = Alignment(horizontal="right")
        ws.cell(r, 3).number_format = ws.cell(r, 4).number_format = ws.cell(r, 5).number_format = "General"
        ok = f"AND(ISNUMBER(C{r}),C{r}>0,ISNUMBER(D{r}),D{r}>0,ISNUMBER(F{r}),F{r}>0)"
        layers = f"MAX(1,INT(IF(AND(ISNUMBER(G{r}),G{r}>0),G{r},1)))"
        ws[f"I{r}"] = f"=IF({ok},ROUNDUP(F{r}/{layers},0),0)"
        ws[f"J{r}"] = f"=IF(I{r}>0,(C{r}/$C$8)*(D{r}/$C$8)/$C$5*I{r},0)"
        ws[f"K{r}"] = f"=IF(I{r}>0,(C{r}/$C$8)*(D{r}/$C$8)*I{r},0)"
        ws[f"L{r}"] = f"=IF(AND(I{r}>0,ISNUMBER(E{r}),E{r}>0),(C{r}/$C$8)*(D{r}/$C$8)*(E{r}/$C$8)*F{r},0)"
        ws[f"M{r}"] = f"=IF(AND(I{r}>0,ISNUMBER(H{r}),H{r}>0),H{r}*F{r},0)"
        ws[f"N{r}"] = (
            f'=IF(COUNTA(B{r}:H{r})=0,"",TRIM('
            f'IF(I{r}=0,"Länge, Breite und Anzahl prüfen. ","")&'
            f'IF(AND(ISNUMBER(D{r}),D{r}/$C$8>$C$5),"Breite größer als Referenzbreite – Fahrzeug prüfen. ","")&'
            f'IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r}),MAX(C{r},D{r})/$C$8>$C$7),"Länger als Vergleichsfahrzeug. ","")&'
            f'IF(AND(ISNUMBER(G{r}),G{r}>1,I{r}>0),"Gestapelt gerechnet – Innenhöhe und Sicherung prüfen. ","")))'
        )
        for col, fmt in (("I", "0;-0;;@"), ("J", "0.0000;-0.0000;;@"), ("K", "0.00;-0.00;;@"), ("L", "0.00;-0.00;;@"), ("M", "#,##0.00;-#,##0.00;;@")):
            cell = ws[f"{col}{r}"]
            cell.number_format = fmt
            cell.fill, cell.font, cell.border = CALC_FILL, F_BASE, BORDER
            cell.alignment = Alignment(horizontal="right")
        ws[f"N{r}"].fill, ws[f"N{r}"].font, ws[f"N{r}"].border = CALC_FILL, F_SMALL, BORDER

    ws.cell(TOTAL, 2, "Summe").font = F_BOLD
    for col, fmt in (("I", "0"), ("J", "0.00"), ("K", "0.00"), ("L", "0.00"), ("M", "#,##0.00")):
        c = ws[f"{col}{TOTAL}"]
        c.value = f"=SUM({col}{FIRST}:{col}{LAST})"
        c.number_format, c.font, c.fill, c.border = fmt, F_BOLD, KEY_FILL, BORDER
        c.alignment = Alignment(horizontal="right")
    for col in "ABCDEFGHN":
        ws[f"{col}{TOTAL}"].border = BORDER
    ws[f"A{TOTAL + 2}"] = (
        "Hinweis: Alle Werte sind rechnerische Planungswerte (Lademeter = Länge × Breite ÷ Referenzbreite × Bodenstellplätze). "
        "Sie ersetzen weder Ladeplan, Achslastprüfung noch Ladungssicherung. Abrechnung und Umrechnungsfaktor richten sich nach dem Tarif Ihrer Spedition."
    )
    ws[f"A{TOTAL + 2}"].font = F_SMALL

    # Bedingte Formatierung
    ws.conditional_formatting.add(f"N{FIRST}:N{LAST}", FormulaRule(formula=[f'LEN(N{FIRST})>0'], fill=WARN_FILL))
    ws.conditional_formatting.add("M7", CellIsRule(operator="lessThan", formula=["0"], fill=WARN_FILL))
    ws.freeze_panes = f"C{FIRST}"

    # Druck
    ws.print_area = f"A1:N{TOTAL + 2}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.oddHeader.center.text = "lademeter-rechner.de – Lademeter-Vorlage"
    ws.oddFooter.center.text = "Seite &P von &N"
    ws.print_title_rows = "13:13"

    # Blattschutz ohne Kennwort (Formeln gegen versehentliches Überschreiben schützen)
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def text_sheet(ws, title, rows, widths=(110,)):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = widths[0]
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    r = 3
    for kind, text in rows:
        cell = ws.cell(r, 1, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "h":
            cell.font = Font(name="Calibri", size=13, bold=True, color="C2410C")
            r += 1
            continue
        if kind == "b":
            cell.font = F_BOLD
        elif kind == "s":
            cell.font = F_SMALL
        elif kind == "l":
            cell.font = F_LINK
            cell.hyperlink = text
        else:
            cell.font = F_BASE
        lines = max(1, int(len(text) / 105) + 1)
        ws.row_dimensions[r].height = 15 * lines
        r += 1


def build():
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Anleitung"
    ws1 = wb.create_sheet("Rechner")
    ws2 = wb.create_sheet("Beispiel")
    ws3 = wb.create_sheet("Referenz")
    ws4 = wb.create_sheet("Hinweise und Quellen")

    ws0.sheet_properties.tabColor = NAVY
    text_sheet(ws0, "Lademeter berechnen – Excel-Vorlage", [
        ("n", f"Version {VERSION} · Stand {STAND} · kostenlos nutzbar (siehe Lizenz auf dem Blatt „Hinweise und Quellen“) · {SITE}"),
        ("h", "So funktioniert die Vorlage"),
        ("n", "1. Blatt „Rechner“ öffnen. Oben die Maßeinheit wählen (mm, cm oder m). Optional den Umrechnungsfaktor Ihrer Spedition in kg je Lademeter eintragen."),
        ("n", "2. Je Ladungsposition eine Zeile ausfüllen: Länge, Breite, Anzahl (Pflicht) sowie Stapellagen, Höhe und Gewicht je Stück (optional). Stapellagen 1 = nicht stapelbar."),
        ("n", "3. Das Ergebnis erscheint oben: Lademeter, Bodenstellplätze, Grundfläche, Volumen, Gewicht sowie – mit Faktor – Lademetergewicht und frachtpflichtiges Gewicht."),
        ("n", "4. Das Blatt „Beispiel“ zeigt eine ausgefüllte Beispielsendung mit demselben Aufbau. Das Blatt „Referenz“ listet die Lademeter gängiger Paletten."),
        ("h", "Die Formeln"),
        ("b", "Bodenstellplätze = Anzahl ÷ Stapellagen, aufgerundet"),
        ("b", "Lademeter je Position = Länge [m] × Breite [m] ÷ 2,40 m × Bodenstellplätze"),
        ("b", "Lademetergewicht = Lademeter × Faktor [kg/ldm]"),
        ("b", "Frachtpflichtiges Gewicht = höherer Wert aus tatsächlichem Gewicht und Lademetergewicht"),
        ("n", "Beispiel: Eine Europalette (1,20 m × 0,80 m) belegt 1,20 × 0,80 ÷ 2,40 = 0,40 Lademeter. Zehn Europaletten nebeneinander ohne Stapeln sind 4,0 Lademeter; zweilagig gestapelt (5 Bodenstellplätze) sind es 2,0 Lademeter."),
        ("h", "Wichtige Hinweise"),
        ("n", "• Die Referenzbreite von 2,40 m ist eine in der Branche übliche Planungsgröße, keine Norm. Der Wert lässt sich ändern, wenn Ihre Spedition etwas anderes vorgibt."),
        ("n", "• Gestapelte Positionen zählen nur mit ihren Bodenstellplätzen. Prüfen Sie selbst, ob Innenhöhe, Belastbarkeit und Ladungssicherung die Lagenzahl erlauben und ob Ihre Spedition gestapelte Ware so abrechnet."),
        ("n", "• Der Umrechnungsfaktor ist eine Tarifregel der jeweiligen Spedition (häufig genannt werden 1.500 bis 2.000 kg je Lademeter). Im Beispiel wird 1.850 kg verwendet; das ist ein Beispielwert, kein Standard."),
        ("n", "• Das Ergebnis ist eine rechnerische Planungshilfe. Es ersetzt keinen Ladeplan, keine Achslastprüfung und keine Ladungssicherung. Speditionen runden Lademeter für die Abrechnung teilweise auf; die Vorlage zeigt den genauen Wert."),
        ("n", "• Zellen mit Formeln sind geschützt (ohne Kennwort: Überprüfen > Blattschutz aufheben), damit sie nicht versehentlich überschrieben werden."),
        ("n", "• Dezimalzahlen: Verwenden Sie das Dezimalzeichen Ihrer Excel-Spracheinstellung (in Deutschland das Komma)."),
    ])

    build_calc_sheet(ws1, example=False)
    build_calc_sheet(ws2, example=True)

    # Referenz
    ws3.sheet_view.showGridLines = False
    ws3.sheet_properties.tabColor = "64748B"
    for col, w in zip("ABCDEFG", (34, 13, 13, 17, 17, 19, 3)):
        ws3.column_dimensions[col].width = w
    ws3["A1"] = "Referenz: Lademeter gängiger Ladeeinheiten"
    ws3["A1"].font = F_TITLE
    ws3["A2"] = "Maße in mm (blau = änderbar). Rechenweg: Länge × Breite ÷ Referenzbreite (Rechner!C5, Standard 2,40 m)."
    ws3["A2"].font = F_SMALL
    for i, h in enumerate(["Ladeeinheit", "Länge (mm)", "Breite (mm)", "Grundfläche (m²)", "Lademeter je Stück", "Stück je Lademeter"], start=1):
        c = ws3.cell(4, i, h)
        c.fill, c.font, c.border = HEAD_FILL, F_HEAD, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws3.row_dimensions[4].height = 32
    units = [("Europalette", 1200, 800), ("Industriepalette", 1200, 1000), ("Halbpalette (Düsseldorfer)", 800, 600), ("eigene Ladeeinheit", None, None)]
    for k, (name, l, b) in enumerate(units):
        r = 5 + k
        ws3.cell(r, 1, name).font = F_BASE
        for col, val in ((2, l), (3, b)):
            c = ws3.cell(r, col, val)
            c.fill, c.font = INPUT_FILL, F_INPUT
            c.protection = Protection(locked=False)
            c.number_format = "#,##0"
        ws3[f"D{r}"] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),B{r}/1000*C{r}/1000,"")'
        ws3[f"E{r}"] = f'=IF(ISNUMBER(D{r}),D{r}/Rechner!$C$5,"")'
        ws3[f"F{r}"] = f'=IF(ISNUMBER(E{r}),1/E{r},"")'
        for col, fmt in (("D", "0.00"), ("E", "0.000"), ("F", "0.00")):
            ws3[f"{col}{r}"].number_format = fmt
            ws3[f"{col}{r}"].fill = CALC_FILL
        for col in "ABCDEF":
            ws3[f"{col}{r}"].border = BORDER
    ws3["A10"] = "Lademeter für mehrere Ladeeinheiten (nicht gestapelt)"
    ws3["A10"].font = F_BOLD
    for i, h in enumerate(["Anzahl", "Europaletten (ldm)", "Industriepaletten (ldm)", "Halbpaletten (ldm)"], start=1):
        c = ws3.cell(11, i, h)
        c.fill, c.font, c.border = HEAD_FILL, F_HEAD, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws3.row_dimensions[11].height = 32
    for k, n in enumerate([1, 2, 4, 5, 8, 10, 15, 20, 26, 33, 34]):
        r = 12 + k
        ws3.cell(r, 1, n).alignment = Alignment(horizontal="center")
        ws3[f"B{r}"] = f"=ROUND($A{r}*$E$5,2)"
        ws3[f"C{r}"] = f"=ROUND($A{r}*$E$6,2)"
        ws3[f"D{r}"] = f"=ROUND($A{r}*$E$7,2)"
        for col in "ABCD":
            ws3[f"{col}{r}"].border = BORDER
            if col != "A":
                ws3[f"{col}{r}"].number_format = "0.00"
                ws3[f"{col}{r}"].fill = CALC_FILL
    ws3.protection.sheet = True

    text_sheet(ws4, "Hinweise, Quellen und Lizenz", [
        ("h", "Methode"),
        ("n", "Die Vorlage bildet die Rechenkerne des Online-Rechners auf lademeter-rechner.de nach. Der Rechenweg (Referenzbreite 2,40 m, Bodenstellplätze, Lademetergewicht, frachtpflichtiges Gewicht) ist auf der Seite „Berechnungsgrundlagen“ dokumentiert."),
        ("l", f"{SITE}/methodik/"),
        ("n", "Es gibt keine Norm, die die Lademeter-Formel festlegt. Es handelt sich um eine Branchenkonvention, die Speditionen unterschiedlich anwenden. Maßgeblich sind immer die Angaben und Tarife der beauftragten Spedition."),
        ("h", "Verifikation"),
        ("n", f"Die Formeln dieser Vorlage wurden am {STAND} mit automatisierten Testfällen gegen die Rechenkerne des Online-Rechners abgeglichen (verschiedene Einheiten, Stapellagen, unvollständige Positionen, Faktor mit und ohne Gewicht). Die Ergebnisse stimmen überein. Dies ist eine technische Prüfung der Rechnung, keine Haftungszusage für den Einzelfall."),
        ("h", "Weiterführende Quellen (Begriffserklärungen)"),
        ("l", "https://de.wikipedia.org/wiki/Lademeter"),
        ("l", "https://www.chemie.de/lexikon/Lademeter.html"),
        ("h", "Haftung"),
        ("n", "Alle Angaben ohne Gewähr. Die Vorlage dient der ersten Orientierung und ersetzt keine verbindliche Ladeplanung inklusive Ladungssicherung, Achslastprüfung und rechtlicher Vorgaben."),
        ("h", "Lizenz und Weitergabe"),
        ("n", f"© {datetime.date.today().year} lademeter-rechner.de. Nutzung, Weitergabe und Anpassung sind erlaubt (CC BY 4.0), wenn die Quelle „lademeter-rechner.de“ genannt wird. Lizenztext: https://creativecommons.org/licenses/by/4.0/deed.de"),
        ("h", "Änderungsprotokoll"),
        ("n", f"Version {VERSION} ({STAND}): erste Veröffentlichung."),
    ])

    wb.properties.title = "Lademeter berechnen – Excel-Vorlage"
    wb.properties.creator = "lademeter-rechner.de"
    wb.properties.description = "Kostenlose Excel-Vorlage zur Berechnung von Lademetern, Bodenstellplätzen und frachtpflichtigem Gewicht."
    wb.properties.keywords = "Lademeter, Excel, Vorlage, Palette, LKW"
    wb.active = 1
    return wb


if __name__ == "__main__":
    import os
    os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
    build().save(OUT)
    print("gespeichert:", OUT)
