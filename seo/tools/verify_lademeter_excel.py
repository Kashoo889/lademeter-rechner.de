"""Gleicht die Excel-Vorlage mit den Rechenkernen von lademeter-rechner.de ab.

Aufruf (im Projektordner):  python seo/tools/verify_lademeter_excel.py [Anzahl Zufallsfälle]
Voraussetzungen: pip install openpyxl formulas ; Node >= 22.18 (führt lib/calculators/lademeter.ts direkt aus)

Ablauf: Testfälle erzeugen -> Rechenkern (Node) -> Excel-Formeln (Python-Bibliothek "formulas") -> Vergleich.
Beendet mit Exit-Code 1, wenn ein Wert abweicht.
"""
import json
import os
import random
import subprocess
import sys
import tempfile
import warnings

import formulas
from openpyxl import load_workbook

warnings.filterwarnings("ignore")
XLSX = os.environ.get("XLSX_PATH", "seo/assets/Lademeter-Berechnung-Excel-Vorlage.xlsx")
N_RANDOM = int(sys.argv[1]) if len(sys.argv) > 1 else 40
FIRST = 14


def fixed_cases():
    return [
        # 10 Europaletten, nicht gestapelt
        {"unit": "cm", "factor": None, "rows": [{"length": 120, "width": 80, "quantity": 10}]},
        # zweilagig
        {"unit": "cm", "factor": None, "rows": [{"length": 120, "width": 80, "quantity": 10, "layers": 2}]},
        # ungerade Anzahl, zweilagig (aufrunden)
        {"unit": "cm", "factor": None, "rows": [{"length": 120, "width": 100, "quantity": 7, "layers": 2}]},
        # Gewicht + Faktor, Lademetergewicht groesser
        {"unit": "cm", "factor": 1850, "rows": [{"length": 120, "width": 80, "quantity": 10, "weight": 420, "height": 150}]},
        # schwere Ware: tatsaechliches Gewicht groesser
        {"unit": "cm", "factor": 1850, "rows": [{"length": 120, "width": 100, "quantity": 6, "weight": 1100}]},
        # Faktor ohne Gewicht
        {"unit": "cm", "factor": 1500, "rows": [{"length": 120, "width": 80, "quantity": 4}]},
        # mm und m
        {"unit": "mm", "factor": None, "rows": [{"length": 1200, "width": 800, "quantity": 33}]},
        {"unit": "m", "factor": 2000, "rows": [{"length": 1.2, "width": 0.8, "quantity": 34, "weight": 300}]},
        # ungueltige Zeile wird ignoriert
        {"unit": "cm", "factor": None, "rows": [{"length": 120, "width": None, "quantity": 5}, {"length": 120, "width": 80, "quantity": 2}]},
        # nur ungueltige Zeilen -> keine Berechnung
        {"unit": "cm", "factor": 1850, "rows": [{"length": None, "width": 80, "quantity": 5}]},
        # Mischladung mit Hoehe und Gewicht teilweise
        {"unit": "cm", "factor": 1850, "rows": [
            {"length": 120, "width": 80, "quantity": 8, "height": 150, "weight": 420},
            {"length": 120, "width": 100, "quantity": 6, "layers": 2},
            {"length": 85, "width": 65, "quantity": 20, "height": 90, "weight": 85}]},
        # breiter als Referenzbreite
        {"unit": "cm", "factor": None, "rows": [{"length": 300, "width": 250, "quantity": 1}]},
    ]


def random_cases(n):
    rng = random.Random(20261020)
    out = []
    for _ in range(n):
        unit = rng.choice(["mm", "cm", "m"])
        scale = {"mm": 1000, "cm": 100, "m": 1}[unit]
        rows = []
        for _ in range(rng.randint(1, 8)):
            row = {
                "length": round(rng.uniform(0.3, 3.0) * scale, 2) if rng.random() > 0.08 else None,
                "width": round(rng.uniform(0.3, 2.6) * scale, 2) if rng.random() > 0.08 else None,
                "quantity": rng.randint(1, 60) if rng.random() > 0.05 else None,
            }
            if rng.random() < 0.6: row["height"] = round(rng.uniform(0.2, 2.5) * scale, 2)
            if rng.random() < 0.6: row["weight"] = round(rng.uniform(5, 900), 1)
            if rng.random() < 0.5: row["layers"] = rng.randint(1, 5)
            rows.append(row)
        out.append({"unit": unit, "factor": rng.choice([None, 1500, 1750, 1850, 2000]), "rows": rows})
    return out


def core_expected(cases):
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as f:
        json.dump(cases, f)
        path = f.name
    res = subprocess.run(["node", "seo/tools/core_expected.mjs", path], capture_output=True, text=True, encoding="utf-8")
    os.unlink(path)
    if res.returncode != 0:
        print(res.stderr)
        sys.exit(2)
    return json.loads(res.stdout)


def excel_values(case):
    wb = load_workbook(XLSX)
    ws = wb["Rechner"]
    ws["C4"] = case["unit"]
    ws["C6"] = case["factor"]
    for i, row in enumerate(case["rows"]):
        r = FIRST + i
        for col, key in (("C", "length"), ("D", "width"), ("E", "height"), ("F", "quantity"), ("G", "layers"), ("H", "weight")):
            ws[f"{col}{r}"] = row.get(key)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "case.xlsx")
        wb.save(path)
        model = formulas.ExcelModel().loads(path).finish()
        sol = model.calculate()
        name = "'[case.xlsx]RECHNER'!"

        def get(cell):
            v = sol[name + cell].value[0, 0]
            return v

        return {c: get(c) for c in ("H4", "H5", "H6", "H7", "H8", "M4", "M5", "M6", "M7")}


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def compare(case, exp, got, idx):
    problems = []

    def chk(label, e, g, tol=0.0051):
        gn = num(g)
        if e is None:
            if gn is not None:
                problems.append(f"{label}: erwartet –, Excel {g}")
        elif gn is None or abs(gn - e) > tol:
            problems.append(f"{label}: erwartet {e}, Excel {g}")

    if exp is None:  # Rechenkern liefert null = nichts zu berechnen
        for cell in ("H4", "H5", "H6", "H7", "H8"):
            chk(cell, 0.0, got[cell])
        for cell in ("M4", "M5", "M6", "M7"):
            chk(cell, None, got[cell])
        return problems
    chk("Lademeter", exp["totalLoadingMeters"], got["H4"])
    chk("Bodenstellplätze", exp["totalFloorPositions"], got["H5"], 0.001)
    chk("Grundfläche", exp["totalFootprintSqm"], got["H6"])
    chk("Volumen", exp["totalVolumeCbm"], got["H7"])
    chk("Gewicht", exp["totalWeightKg"], got["H8"], 0.011)
    chk("Lademetergewicht", exp["loadingMeterWeightKg"], got["M4"], 0.011)
    chk("frachtpflichtiges Gewicht", exp["chargeableWeightKg"], got["M5"], 0.011)
    chk("Anteil Sattelzug %", exp["semiTrailerSharePercent"], got["M6"], 0.011)
    chk("Rest-Lademeter", round(13.6 - exp["totalLoadingMeters"], 2), got["M7"], 0.011)
    return problems


def main():
    cases = fixed_cases() + random_cases(N_RANDOM)
    expected = core_expected(cases)
    failures = 0
    for i, (case, exp) in enumerate(zip(cases, expected)):
        got = excel_values(case)
        problems = compare(case, exp, got, i)
        if problems:
            failures += 1
            print(f"FEHLER Fall {i}: {json.dumps(case, ensure_ascii=False)}")
            for p in problems:
                print("   ", p)
    print(f"{len(cases)} Fälle geprüft ({len(fixed_cases())} feste + {N_RANDOM} zufällige), Abweichungen: {failures}")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
